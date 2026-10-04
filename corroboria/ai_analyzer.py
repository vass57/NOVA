# ============================================================
# CORROBORIA
# AI / MACHINE LEARNING ANALYSIS
# ============================================================
#
# PURPOSE
#
# AI is used for:
#
#   1. Ambiguous assignment matching
#   2. Anomaly prioritization
#   3. Recurring-pattern detection
#
#
# IMPORTANT PRINCIPLE
#
# Deterministic rules remain authoritative.
#
# AI DOES NOT OVERRIDE:
#
#   CONFORME
#   ECART_JUSTIFIE
#   ANOMALIE
#
#
# AMBIGUOUS ASSIGNMENT APPROACH
#
# We now train a supervised classifier using:
#
#   POSITIVE EXAMPLES
#       Known deterministic Source ↔ Destination matches
#
#   NEGATIVE EXAMPLES
#       Deliberately incorrect cross-employee pairings
#
# Employee IDs themselves are NOT model features.
#
# Then, for an ambiguous employee, we:
#
#   1. Score every possible Source ↔ Destination pair
#   2. Enumerate complete one-to-one assignments
#   3. Score each complete matching
#   4. Compare best vs second-best matching
#   5. Recommend only when separation is sufficient
#
#
# This prevents the old behavior where four tied candidate
# scores caused idxmax() to arbitrarily choose the first row.
#
# ============================================================


# ============================================================
# IMPORTS
# ============================================================

from itertools import permutations
from pathlib import Path
import ast
import math
import random

import numpy as np
import pandas as pd

from sklearn.ensemble import (
    IsolationForest,
    RandomForestClassifier,
)

from sklearn.impute import SimpleImputer
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler

from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)

from openpyxl.utils import (
    get_column_letter,
)


# ============================================================
# GENERAL HELPERS
# ============================================================

def is_missing(value):
    """
    Safe missing-value test.
    """

    if value is None:
        return True


    try:

        return bool(
            pd.isna(
                value
            )
        )

    except (
        TypeError,
        ValueError,
    ):

        return False


def text_value(value):
    """
    Convert comparable values to stripped text.
    """

    if is_missing(
        value
    ):

        return None


    return str(
        value
    ).strip()


def number_value(value):
    """
    Safely convert a value to float.
    """

    if is_missing(
        value
    ):

        return None


    try:

        return float(
            value
        )

    except (
        TypeError,
        ValueError,
    ):

        return None


def neutral_match(
    left,
    right,
):
    """
    Return:

        1.0 = same value
        0.0 = different values
        0.5 = comparison unavailable because something missing

    Missing information is represented separately using
    missingness features.
    """

    left = text_value(
        left
    )

    right = text_value(
        right
    )


    if (
        left is None
        or right is None
    ):

        return 0.5


    return float(
        left == right
    )


def numeric_gap(
    left,
    right,
):
    """
    Absolute numeric difference.
    """

    left_number = number_value(
        left
    )

    right_number = number_value(
        right
    )


    if (
        left_number is None
        or right_number is None
    ):

        return np.nan


    return abs(
        left_number
        -
        right_number
    )


def missing_mismatch(
    left,
    right,
):
    """
    Return 1 if exactly one side is missing.
    """

    return float(
        is_missing(
            left
        )
        !=
        is_missing(
            right
        )
    )


def ensure_list(value):
    """
    Candidate-index fields may already be lists or may be
    serialized strings.

    Convert either form to a normal Python list.
    """

    if isinstance(
        value,
        list,
    ):

        return value


    if isinstance(
        value,
        tuple,
    ):

        return list(
            value
        )


    if is_missing(
        value
    ):

        return []


    if isinstance(
        value,
        str,
    ):

        try:

            parsed = ast.literal_eval(
                value
            )


            if isinstance(
                parsed,
                (
                    list,
                    tuple,
                ),
            ):

                return list(
                    parsed
                )

        except (
            ValueError,
            SyntaxError,
        ):

            pass


    return [
        value
    ]


# ============================================================
# DESTINATION ASSIGNMENT TYPE
# ============================================================

def destination_assignment_type(
    destination_row,
):
    """
    Reconstruct P / A / S from System B flags.
    """

    primary = destination_row.get(
        "isPrimaryAssignment"
    )

    temporary = destination_row.get(
        "isTemporaryAssignment"
    )


    if (
        is_missing(
            primary
        )
        or
        is_missing(
            temporary
        )
    ):

        return None


    primary = bool(
        primary
    )

    temporary = bool(
        temporary
    )


    if (
        primary
        and
        not temporary
    ):

        return "P"


    if (
        not primary
        and
        temporary
    ):

        return "A"


    if (
        not primary
        and
        not temporary
    ):

        return "S"


    return None


# ============================================================
# PAIR FEATURE ENGINEERING
# ============================================================

PAIR_FEATURE_COLUMNS = [
    "position_id_match",
    "position_code_match",
    "site_code_match",
    "division_id_match",
    "division_code_match",
    "pay_grade_match",
    "assignment_type_match",
    "weekly_hours_gap",
    "daily_hours_gap",
    "weekly_missing_mismatch",
    "daily_missing_mismatch",
]


def build_pair_features(
    source_row,
    destination_row,
):
    """
    Describe one possible Source ↔ Destination assignment pair.

    Only mapped / corroborated information is used.

    We intentionally avoid:

        employee_id
        positionName
        assignmentStartDate

    employee_id would make synthetic negatives trivial to
    distinguish and would not help ambiguous rows belonging to
    the same employee.

    positionName and assignmentStartDate are excluded because
    our deterministic analysis already identified systematic
    discrepancies in those fields.
    """

    source_assignment = text_value(
        source_row.get(
            "TypeAffectation"
        )
    )


    destination_assignment = (
        destination_assignment_type(
            destination_row
        )
    )


    return {
        # ----------------------------------------------------
        # ROLE / EMPLOYMENT
        # ----------------------------------------------------

        "position_id_match":
            neutral_match(
                source_row.get(
                    "CodeEmploi"
                ),
                destination_row.get(
                    "positionId"
                ),
            ),

        "position_code_match":
            neutral_match(
                source_row.get(
                    "CodeEmploi"
                ),
                destination_row.get(
                    "positionCode"
                ),
            ),

        # ----------------------------------------------------
        # SITE
        # ----------------------------------------------------

        "site_code_match":
            neutral_match(
                source_row.get(
                    "CodeSite"
                ),
                destination_row.get(
                    "siteCode"
                ),
            ),

        # ----------------------------------------------------
        # ADMINISTRATIVE UNIT
        # ----------------------------------------------------

        "division_id_match":
            neutral_match(
                source_row.get(
                    "CodeDirection"
                ),
                destination_row.get(
                    "divisionId"
                ),
            ),

        "division_code_match":
            neutral_match(
                source_row.get(
                    "CodeImputation"
                ),
                destination_row.get(
                    "divisionCode"
                ),
            ),

        # ----------------------------------------------------
        # PAY GRADE
        # ----------------------------------------------------

        "pay_grade_match":
            neutral_match(
                source_row.get(
                    "ÉchelleSalariale"
                ),
                destination_row.get(
                    "payGradeId"
                ),
            ),

        # ----------------------------------------------------
        # ASSIGNMENT TYPE
        # ----------------------------------------------------

        "assignment_type_match":
            neutral_match(
                source_assignment,
                destination_assignment,
            ),

        # ----------------------------------------------------
        # HOURS
        # ----------------------------------------------------

        "weekly_hours_gap":
            numeric_gap(
                source_row.get(
                    "HeuresNormeHebdo"
                ),
                destination_row.get(
                    "weeklyHoursOverride"
                ),
            ),

        "daily_hours_gap":
            numeric_gap(
                source_row.get(
                    "HeuresNormeQuotidienne"
                ),
                destination_row.get(
                    "dailyHoursOverride"
                ),
            ),

        "weekly_missing_mismatch":
            missing_mismatch(
                source_row.get(
                    "HeuresNormeHebdo"
                ),
                destination_row.get(
                    "weeklyHoursOverride"
                ),
            ),

        "daily_missing_mismatch":
            missing_mismatch(
                source_row.get(
                    "HeuresNormeQuotidienne"
                ),
                destination_row.get(
                    "dailyHoursOverride"
                ),
            ),
    }


# ============================================================
# KNOWN POSITIVE PAIRS
# ============================================================

def get_deterministic_pairs(
    assignment_matches_df,
):
    """
    Return safely matched assignments.
    """

    return assignment_matches_df[

        assignment_matches_df[
            "assignment_match_status"
        ]
        ==
        "MATCHED"

    ].copy()


# ============================================================
# TRAINING DATA
# ============================================================

def build_pair_training_data(
    source_df,
    destination_df,
    assignment_matches_df,
    negative_ratio=5,
    random_state=42,
):
    """
    Build classifier training data.

    POSITIVES
    ---------

    Every deterministic assignment match.

    NEGATIVES
    ---------

    Pair a deterministic source assignment with destination
    assignments belonging to OTHER employees.

    IDs themselves are never used as model features.

    We generate up to:

        positive_count × negative_ratio

    negative observations.
    """

    deterministic = get_deterministic_pairs(
        assignment_matches_df
    )


    training_rows = []


    # ========================================================
    # POSITIVE EXAMPLES
    # ========================================================

    for _, match in deterministic.iterrows():

        source_index = int(
            match[
                "source_row_index"
            ]
        )


        destination_index = int(
            match[
                "destination_row_index"
            ]
        )


        features = build_pair_features(

            source_row=
                source_df.loc[
                    source_index
                ],

            destination_row=
                destination_df.loc[
                    destination_index
                ],
        )


        training_rows.append(
            {
                **features,

                "label":
                    1,

                "example_type":
                    "KNOWN_MATCH",

                "source_row_index":
                    source_index,

                "destination_row_index":
                    destination_index,
            }
        )


    # ========================================================
    # NEGATIVE POOL
    # ========================================================

    negative_pool = []


    deterministic_records = (
        deterministic
        .to_dict(
            orient="records"
        )
    )


    for source_match in deterministic_records:

        source_index = int(
            source_match[
                "source_row_index"
            ]
        )


        source_employee = str(
            source_match[
                "employee_id"
            ]
        )


        source_row = source_df.loc[
            source_index
        ]


        for destination_match in deterministic_records:

            destination_employee = str(
                destination_match[
                    "employee_id"
                ]
            )


            # ------------------------------------------------
            # Only deliberate cross-employee false pairs
            # ------------------------------------------------

            if (
                source_employee
                ==
                destination_employee
            ):

                continue


            destination_index = int(
                destination_match[
                    "destination_row_index"
                ]
            )


            destination_row = destination_df.loc[
                destination_index
            ]


            features = build_pair_features(

                source_row=
                    source_row,

                destination_row=
                    destination_row,
            )


            negative_pool.append(
                {
                    **features,

                    "label":
                        0,

                    "example_type":
                        "SYNTHETIC_NON_MATCH",

                    "source_row_index":
                        source_index,

                    "destination_row_index":
                        destination_index,
                }
            )


    # ========================================================
    # SAMPLE NEGATIVES
    # ========================================================

    rng = random.Random(
        random_state
    )


    positive_count = len(
        training_rows
    )


    desired_negative_count = min(
        len(
            negative_pool
        ),
        positive_count
        *
        negative_ratio,
    )


    if (
        desired_negative_count
        <
        len(
            negative_pool
        )
    ):

        negative_rows = rng.sample(
            negative_pool,
            desired_negative_count,
        )

    else:

        negative_rows = (
            negative_pool
        )


    training_rows.extend(
        negative_rows
    )


    training_df = pd.DataFrame(
        training_rows
    )


    return training_df


# ============================================================
# CLASSIFIER
# ============================================================

def build_pair_classifier():
    """
    Classifier for Source ↔ Destination assignment similarity.

    Random Forest works well here because:

        - feature count is small
        - interactions matter
        - data is tiny
        - nonlinear combinations are possible
        - probabilities/scores are easy to expose
    """

    return Pipeline(
        steps=[
            (
                "imputer",
                SimpleImputer(
                    strategy="median"
                ),
            ),
            (
                "model",
                RandomForestClassifier(
                    n_estimators=500,
                    random_state=42,
                    class_weight="balanced",
                    min_samples_leaf=1,
                    max_features="sqrt",
                ),
            ),
        ]
    )


# ============================================================
# MODEL INFORMATION
# ============================================================

def get_positive_class_index(
    classifier,
):
    """
    Find column corresponding to class 1 in predict_proba().
    """

    model = classifier.named_steps[
        "model"
    ]


    classes = list(
        model.classes_
    )


    return classes.index(
        1
    )


# ============================================================
# EVIDENCE EXPLANATION
# ============================================================

def explain_pair_features(
    features,
):
    """
    Transparent summary of evidence used for one candidate.

    This reports MODEL INPUT EVIDENCE.

    It is not hidden chain-of-thought reasoning.
    """

    matched_labels = []


    mismatched_labels = []


    labels = {
        "position_id_match":
            "positionId",

        "position_code_match":
            "positionCode",

        "site_code_match":
            "siteCode",

        "division_id_match":
            "divisionId",

        "division_code_match":
            "divisionCode",

        "pay_grade_match":
            "payGradeId",

        "assignment_type_match":
            "type d'affectation",
    }


    for (
        feature,
        label,
    ) in labels.items():

        value = features.get(
            feature
        )


        if value == 1.0:

            matched_labels.append(
                label
            )


        elif value == 0.0:

            mismatched_labels.append(
                label
            )


    parts = []


    if matched_labels:

        parts.append(
            "Correspondances: "
            +
            ", ".join(
                matched_labels
            )
        )


    if mismatched_labels:

        parts.append(
            "Différences: "
            +
            ", ".join(
                mismatched_labels
            )
        )


    weekly_gap = features.get(
        "weekly_hours_gap"
    )


    if not pd.isna(
        weekly_gap
    ):

        parts.append(
            f"écart heures/semaine={weekly_gap:g}"
        )


    daily_gap = features.get(
        "daily_hours_gap"
    )


    if not pd.isna(
        daily_gap
    ):

        parts.append(
            f"écart heures/jour={daily_gap:g}"
        )


    if not parts:

        return (
            "Peu d'information comparable disponible."
        )


    return (
        ". ".join(
            parts
        )
        +
        "."
    )


# ============================================================
# COMPLETE MATCHING ENUMERATION
# ============================================================

def enumerate_complete_matchings(
    candidate_df,
    source_candidates,
    destination_candidates,
):
    """
    Enumerate every possible one-to-one assignment.

    Example:

        sources:
            21, 22

        destinations:
            19, 21

    gives:

        Matching A:
            21 -> 19
            22 -> 21

        Matching B:
            21 -> 21
            22 -> 19

    Matching score is the mean pair probability.
    """

    if (
        len(
            source_candidates
        )
        !=
        len(
            destination_candidates
        )
    ):

        return []


    # Prevent combinatorial explosion on unexpected data.
    if len(
        source_candidates
    ) > 7:

        return []


    matching_results = []


    matching_number = 0


    for destination_order in permutations(
        destination_candidates
    ):

        matching_number += 1


        pair_records = []


        probabilities = []


        for (
            source_index,
            destination_index,
        ) in zip(
            source_candidates,
            destination_order,
        ):

            row = candidate_df[
                (
                    candidate_df[
                        "source_row_index"
                    ]
                    ==
                    int(
                        source_index
                    )
                )
                &
                (
                    candidate_df[
                        "destination_row_index"
                    ]
                    ==
                    int(
                        destination_index
                    )
                )
            ]


            if row.empty:

                continue


            pair_probability = float(
                row.iloc[
                    0
                ][
                    "pair_match_score"
                ]
            )


            probabilities.append(
                pair_probability
            )


            pair_records.append(
                (
                    int(
                        source_index
                    ),
                    int(
                        destination_index
                    ),
                )
            )


        if (
            len(
                probabilities
            )
            !=
            len(
                source_candidates
            )
        ):

            continue


        matching_score = float(
            np.mean(
                probabilities
            )
        )


        # Geometric mean gives another useful view and strongly
        # penalizes a matching containing one very poor pair.
        clipped = np.clip(
            probabilities,
            1e-6,
            1.0,
        )


        geometric_score = float(
            np.exp(
                np.mean(
                    np.log(
                        clipped
                    )
                )
            )
        )


        matching_results.append(
            {
                "matching_id":
                    matching_number,

                "pairs":
                    pair_records,

                "matching_score":
                    matching_score,

                "matching_geometric_score":
                    geometric_score,
            }
        )


    matching_results = sorted(
        matching_results,
        key=lambda item: (
            item[
                "matching_score"
            ],
            item[
                "matching_geometric_score"
            ],
        ),
        reverse=True,
    )


    return matching_results


# ============================================================
# MATCHING CONFIDENCE
# ============================================================

def calculate_matching_confidence(
    best_score,
    second_score,
):
    """
    Measure separation between best and second-best complete
    assignment.

    Equal scores:
        0.5 confidence

    Large separation:
        approaches 1.0
    """

    if second_score is None:

        return 1.0


    total = (
        best_score
        +
        second_score
    )


    if total <= 1e-9:

        return 0.5


    confidence = (
        best_score
        /
        total
    )


    return round(
        max(
            0.5,
            min(
                1.0,
                float(
                    confidence
                ),
            ),
        ),
        3,
    )


# ============================================================
# AMBIGUOUS ASSIGNMENT ANALYSIS
# ============================================================

def analyze_ambiguous_assignments(
    source_df,
    destination_df,
    assignment_matches_df,
):
    """
    Analyze ambiguous assignment groups.

    Unlike the previous IsolationForest approach, this method
    explicitly learns MATCH vs NON-MATCH behavior and compares
    complete one-to-one assignment possibilities.
    """

    # ========================================================
    # TRAINING DATA
    # ========================================================

    training_df = (
        build_pair_training_data(

            source_df=
                source_df,

            destination_df=
                destination_df,

            assignment_matches_df=
                assignment_matches_df,
        )
    )


    if training_df.empty:

        return pd.DataFrame()


    if training_df[
        "label"
    ].nunique() < 2:

        return pd.DataFrame()


    positive_count = int(
        (
            training_df[
                "label"
            ]
            ==
            1
        ).sum()
    )


    negative_count = int(
        (
            training_df[
                "label"
            ]
            ==
            0
        ).sum()
    )


    # ========================================================
    # TRAIN CLASSIFIER
    # ========================================================

    classifier = (
        build_pair_classifier()
    )


    classifier.fit(
        training_df[
            PAIR_FEATURE_COLUMNS
        ],
        training_df[
            "label"
        ],
    )


    positive_class_index = (
        get_positive_class_index(
            classifier
        )
    )


    # ========================================================
    # AMBIGUOUS GROUPS
    # ========================================================

    ambiguous = assignment_matches_df[

        assignment_matches_df[
            "assignment_match_status"
        ]
        ==
        "AMBIGUOUS_ASSIGNMENT_MATCH"

    ]


    output_rows = []


    for _, group in ambiguous.iterrows():

        source_candidates = [
            int(
                value
            )
            for value
            in ensure_list(
                group.get(
                    "source_candidate_indices"
                )
            )
        ]


        destination_candidates = [
            int(
                value
            )
            for value
            in ensure_list(
                group.get(
                    "destination_candidate_indices"
                )
            )
        ]


        candidate_rows = []


        # ====================================================
        # SCORE EVERY POSSIBLE PAIR
        # ====================================================

        for source_index in source_candidates:

            source_row = source_df.loc[
                source_index
            ]


            for destination_index in destination_candidates:

                destination_row = destination_df.loc[
                    destination_index
                ]


                features = build_pair_features(

                    source_row=
                        source_row,

                    destination_row=
                        destination_row,
                )


                candidate_rows.append(
                    {
                        "employee_id":
                            group[
                                "employee_id"
                            ],

                        "assignment_type":
                            group[
                                "assignment_type"
                            ],

                        "source_row_index":
                            source_index,

                        "destination_row_index":
                            destination_index,

                        **features,
                    }
                )


        candidate_df = pd.DataFrame(
            candidate_rows
        )


        if candidate_df.empty:

            continue


        probabilities = (
            classifier.predict_proba(
                candidate_df[
                    PAIR_FEATURE_COLUMNS
                ]
            )[
                :,
                positive_class_index
            ]
        )


        candidate_df[
            "pair_match_score"
        ] = probabilities


        # ====================================================
        # ENUMERATE COMPLETE ONE-TO-ONE MATCHINGS
        # ====================================================

        matching_results = (
            enumerate_complete_matchings(

                candidate_df=
                    candidate_df,

                source_candidates=
                    source_candidates,

                destination_candidates=
                    destination_candidates,
            )
        )


        # ====================================================
        # CANNOT FORM COMPLETE MATCHING
        # ====================================================

        if not matching_results:

            for _, candidate in candidate_df.iterrows():

                feature_dict = {
                    feature:
                        candidate[
                            feature
                        ]
                    for feature
                    in PAIR_FEATURE_COLUMNS
                }


                output_rows.append(
                    {
                        "employee_id":
                            candidate[
                                "employee_id"
                            ],

                        "assignment_type":
                            candidate[
                                "assignment_type"
                            ],

                        "source_row_index":
                            candidate[
                                "source_row_index"
                            ],

                        "destination_row_index":
                            candidate[
                                "destination_row_index"
                            ],

                        "pair_match_score":
                            round(
                                float(
                                    candidate[
                                        "pair_match_score"
                                    ]
                                ),
                                6,
                            ),

                        "selected_by_best_matching":
                            False,

                        "best_matching_score":
                            None,

                        "second_matching_score":
                            None,

                        "matching_confidence":
                            0.5,

                        "matching_margin":
                            None,

                        "recommendation":
                            "REVUE_HUMAINE",

                        "recommended_action":
                            (
                                "Le nombre de candidats source et "
                                "destination ne permet pas une "
                                "affectation biunivoque automatique."
                            ),

                        "evidence_summary":
                            explain_pair_features(
                                feature_dict
                            ),

                        "training_positive_count":
                            positive_count,

                        "training_negative_count":
                            negative_count,

                        "ai_method":
                            (
                                "RandomForestClassifier "
                                "sur paires positives et négatives"
                            ),

                        "deterministic_verdict":
                            "A_INVESTIGUER",
                    }
                )


            continue


        # ====================================================
        # BEST / SECOND MATCHING
        # ====================================================

        best_matching = matching_results[
            0
        ]


        second_matching = (
            matching_results[
                1
            ]
            if len(
                matching_results
            ) > 1
            else
            None
        )


        best_score = float(
            best_matching[
                "matching_score"
            ]
        )


        second_score = (
            float(
                second_matching[
                    "matching_score"
                ]
            )
            if second_matching
            is not None
            else
            None
        )


        margin = (
            best_score
            -
            second_score
            if second_score
            is not None
            else
            best_score
        )


        confidence = (
            calculate_matching_confidence(
                best_score=
                    best_score,

                second_score=
                    second_score,
            )
        )


        # ====================================================
        # TIE DETECTION
        # ====================================================

        tied = (
            second_score
            is not None
            and
            abs(
                best_score
                -
                second_score
            )
            <
            0.01
        )


        # ====================================================
        # RECOMMENDATION POLICY
        # ====================================================

        strong_recommendation = (
            not tied
            and
            best_score >= 0.60
            and
            confidence >= 0.60
        )


        best_pairs = set(
            best_matching[
                "pairs"
            ]
        )


        # ====================================================
        # OUTPUT EACH UNIQUE CANDIDATE PAIR ONCE
        # ====================================================

        for _, candidate in candidate_df.iterrows():

            pair = (
                int(
                    candidate[
                        "source_row_index"
                    ]
                ),
                int(
                    candidate[
                        "destination_row_index"
                    ]
                ),
            )


            selected = (
                pair
                in
                best_pairs
            )


            feature_dict = {
                feature:
                    candidate[
                        feature
                    ]
                for feature
                in PAIR_FEATURE_COLUMNS
            }


            # ------------------------------------------------
            # EXACT / NEAR TIE
            # ------------------------------------------------

            if tied:

                recommendation = (
                    "REVUE_HUMAINE"
                )


                recommended_action = (
                    "Les meilleures combinaisons obtiennent "
                    "des scores trop proches. Aucune paire "
                    "n'est sélectionnée automatiquement."
                )


                selected_for_output = (
                    False
                )


            # ------------------------------------------------
            # STRONG COMPLETE MATCHING
            # ------------------------------------------------

            elif strong_recommendation:

                if selected:

                    recommendation = (
                        "PROPOSITION_IA"
                    )


                    recommended_action = (
                        "Cette paire appartient à la meilleure "
                        "combinaison biunivoque trouvée par le "
                        "modèle. Valider fonctionnellement avant "
                        "acceptation."
                    )


                else:

                    recommendation = (
                        "CANDIDAT_ALTERNATIF"
                    )


                    recommended_action = (
                        "Cette paire n'appartient pas à la "
                        "meilleure combinaison globale."
                    )


                selected_for_output = (
                    selected
                )


            # ------------------------------------------------
            # BEST EXISTS BUT CONFIDENCE TOO LOW
            # ------------------------------------------------

            else:

                recommendation = (
                    "REVUE_HUMAINE"
                    if selected
                    else
                    "CANDIDAT_ALTERNATIF"
                )


                recommended_action = (
                    "Une meilleure combinaison existe, mais "
                    "la séparation ou le score reste "
                    "insuffisant pour une recommandation "
                    "automatique forte."
                )


                selected_for_output = (
                    selected
                )


            output_rows.append(
                {
                    "employee_id":
                        candidate[
                            "employee_id"
                        ],

                    "assignment_type":
                        candidate[
                            "assignment_type"
                        ],

                    "source_row_index":
                        candidate[
                            "source_row_index"
                        ],

                    "destination_row_index":
                        candidate[
                            "destination_row_index"
                        ],

                    "pair_match_score":
                        round(
                            float(
                                candidate[
                                    "pair_match_score"
                                ]
                            ),
                            6,
                        ),

                    "selected_by_best_matching":
                        bool(
                            selected_for_output
                        ),

                    "best_matching_score":
                        round(
                            best_score,
                            6,
                        ),

                    "second_matching_score":
                        (
                            round(
                                second_score,
                                6,
                            )
                            if second_score
                            is not None
                            else
                            None
                        ),

                    "matching_confidence":
                        confidence,

                    "matching_margin":
                        round(
                            float(
                                margin
                            ),
                            6,
                        ),

                    "recommendation":
                        recommendation,

                    "recommended_action":
                        recommended_action,

                    "evidence_summary":
                        explain_pair_features(
                            feature_dict
                        ),

                    "training_positive_count":
                        positive_count,

                    "training_negative_count":
                        negative_count,

                    "ai_method":
                        (
                            "RandomForestClassifier + "
                            "optimisation de correspondance "
                            "biunivoque"
                        ),

                    "deterministic_verdict":
                        "A_INVESTIGUER",
                }
            )


    result_df = pd.DataFrame(
        output_rows
    )


    if not result_df.empty:

        result_df = (
            result_df
            .sort_values(
                by=[
                    "employee_id",
                    "selected_by_best_matching",
                    "pair_match_score",
                ],
                ascending=[
                    True,
                    False,
                    False,
                ],
            )
            .reset_index(
                drop=True
            )
        )


    return result_df


# ============================================================
# EMPLOYEE ANOMALY FEATURES
# ============================================================

ANOMALY_FEATURE_COLUMNS = [
    "anomaly_count",
    "investigate_count",
    "direct_anomaly_count",
    "rule_anomaly_count",
    "structural_anomaly_count",
    "critical_count",
    "high_priority_count",
    "unique_anomaly_fields",
    "mean_field_rarity",
    "max_field_rarity",
]


def build_employee_anomaly_features(
    final_report_df,
):
    """
    Build anomaly signature per employee.
    """

    employees = (
        final_report_df[
            "employee_id"
        ]
        .dropna()
        .unique()
        .tolist()
    )


    anomaly_rows = final_report_df[

        final_report_df[
            "verdict"
        ]
        ==
        "ANOMALIE"

    ].copy()


    total_employees = max(
        len(
            employees
        ),
        1,
    )


    # ========================================================
    # FIELD PREVALENCE
    # ========================================================

    if anomaly_rows.empty:

        field_prevalence = {}


    else:

        field_prevalence = (
            anomaly_rows
            .groupby(
                "field"
            )[
                "employee_id"
            ]
            .nunique()
            .div(
                total_employees
            )
            .to_dict()
        )


    rows = []


    for employee_id in employees:

        employee_rows = final_report_df[

            final_report_df[
                "employee_id"
            ]
            ==
            employee_id

        ]


        anomalies = employee_rows[

            employee_rows[
                "verdict"
            ]
            ==
            "ANOMALIE"

        ]


        investigations = employee_rows[

            employee_rows[
                "verdict"
            ]
            ==
            "A_INVESTIGUER"

        ]


        rarity_values = []


        for field in anomalies[
            "field"
        ].dropna():

            prevalence = (
                field_prevalence.get(
                    field,
                    0.0,
                )
            )


            rarity_values.append(
                1.0
                -
                prevalence
            )


        mean_rarity = (
            float(
                np.mean(
                    rarity_values
                )
            )
            if rarity_values
            else
            0.0
        )


        max_rarity = (
            float(
                np.max(
                    rarity_values
                )
            )
            if rarity_values
            else
            0.0
        )


        rows.append(
            {
                "employee_id":
                    employee_id,

                "anomaly_count":
                    len(
                        anomalies
                    ),

                "investigate_count":
                    len(
                        investigations
                    ),

                "direct_anomaly_count":
                    int(
                        (
                            anomalies[
                                "comparison_layer"
                            ]
                            ==
                            "DIRECT"
                        ).sum()
                    ),

                "rule_anomaly_count":
                    int(
                        (
                            anomalies[
                                "comparison_layer"
                            ]
                            ==
                            "BUSINESS_RULE"
                        ).sum()
                    ),

                "structural_anomaly_count":
                    int(
                        (
                            anomalies[
                                "comparison_layer"
                            ]
                            ==
                            "STRUCTURAL"
                        ).sum()
                    ),

                "critical_count":
                    int(
                        (
                            anomalies[
                                "priority"
                            ]
                            ==
                            "CRITIQUE"
                        ).sum()
                    ),

                "high_priority_count":
                    int(
                        (
                            anomalies[
                                "priority"
                            ]
                            ==
                            "ÉLEVÉE"
                        ).sum()
                    ),

                "unique_anomaly_fields":
                    anomalies[
                        "field"
                    ].nunique(),

                "mean_field_rarity":
                    mean_rarity,

                "max_field_rarity":
                    max_rarity,
            }
        )


    return pd.DataFrame(
        rows
    )


# ============================================================
# AI PRIORITIZATION
# ============================================================

def prioritize_anomalies(
    final_report_df,
):
    """
    Rank employees by unusual anomaly profile.

    This adjusts INVESTIGATION PRIORITY only.

    It never changes deterministic verdicts.
    """

    feature_df = (
        build_employee_anomaly_features(
            final_report_df=
                final_report_df
        )
    )


    if len(
        feature_df
    ) < 5:

        return pd.DataFrame()


    model = Pipeline(
        steps=[
            (
                "imputer",
                SimpleImputer(
                    strategy="median"
                ),
            ),
            (
                "scaler",
                StandardScaler(),
            ),
            (
                "model",
                IsolationForest(
                    n_estimators=500,
                    contamination="auto",
                    random_state=42,
                ),
            ),
        ]
    )


    model.fit(
        feature_df[
            ANOMALY_FEATURE_COLUMNS
        ]
    )


    raw_scores = (
        -model.score_samples(
            feature_df[
                ANOMALY_FEATURE_COLUMNS
            ]
        )
    )


    minimum = float(
        raw_scores.min()
    )


    maximum = float(
        raw_scores.max()
    )


    if abs(
        maximum
        -
        minimum
    ) < 1e-9:

        normalized_scores = (
            np.zeros(
                len(
                    raw_scores
                )
            )
        )


    else:

        normalized_scores = (
            raw_scores
            -
            minimum
        ) / (
            maximum
            -
            minimum
        )


    feature_df[
        "ai_anomaly_score"
    ] = normalized_scores


    priorities = []


    explanations = []


    for _, row in feature_df.iterrows():

        score = float(
            row[
                "ai_anomaly_score"
            ]
        )


        if row[
            "critical_count"
        ] > 0:

            priority = (
                "CRITIQUE"
            )


        elif score >= 0.75:

            priority = (
                "ÉLEVÉE"
            )


        elif score >= 0.40:

            priority = (
                "MOYENNE"
            )


        else:

            priority = (
                "FAIBLE"
            )


        priorities.append(
            priority
        )


        reasons = []


        if row[
            "structural_anomaly_count"
        ] > 0:

            reasons.append(
                "anomalie structurelle"
            )


        if row[
            "direct_anomaly_count"
        ] > 0:

            reasons.append(
                (
                    f"{int(row['direct_anomaly_count'])} "
                    f"écart(s) direct(s)"
                )
            )


        if row[
            "mean_field_rarity"
        ] >= 0.5:

            reasons.append(
                "champs d'anomalie relativement rares"
            )


        if row[
            "anomaly_count"
        ] >= 4:

            reasons.append(
                "plusieurs anomalies simultanées"
            )


        if not reasons:

            reasons.append(
                "profil proche des anomalies récurrentes"
            )


        explanations.append(
            (
                f"Score IA={score:.3f}. "
                +
                "; ".join(
                    reasons
                )
                +
                "."
            )
        )


    feature_df[
        "ai_priority"
    ] = priorities


    feature_df[
        "ai_explanation"
    ] = explanations


    feature_df = feature_df[

        (
            feature_df[
                "anomaly_count"
            ]
            >
            0
        )
        |
        (
            feature_df[
                "investigate_count"
            ]
            >
            0
        )

    ].copy()


    feature_df = (
        feature_df
        .sort_values(
            by=[
                "ai_anomaly_score",
                "anomaly_count",
            ],
            ascending=[
                False,
                False,
            ],
        )
        .reset_index(
            drop=True
        )
    )


    return feature_df


# ============================================================
# PATTERN ANALYSIS
# ============================================================

def summarize_anomaly_patterns(
    final_report_df,
):
    """
    Distinguish:

        SYSTEMATIQUE
        RECURRENT
        ISOLE

    using field prevalence across employees.
    """

    interesting = final_report_df[

        final_report_df[
            "verdict"
        ].isin(
            [
                "ANOMALIE",
                "A_INVESTIGUER",
            ]
        )

    ].copy()


    if interesting.empty:

        return pd.DataFrame()


    total_employees = max(
        final_report_df[
            "employee_id"
        ]
        .nunique(),
        1,
    )


    patterns = (
        interesting
        .groupby(
            [
                "field",
                "rule_id",
                "verdict",
                "decision_source",
            ],
            dropna=False,
        )
        .agg(
            occurrences=(
                "employee_id",
                "size",
            ),

            affected_employees=(
                "employee_id",
                "nunique",
            ),
        )
        .reset_index()
    )


    patterns[
        "employee_prevalence"
    ] = (

        patterns[
            "affected_employees"
        ]
        /
        total_employees

    )


    classifications = []


    for prevalence in patterns[
        "employee_prevalence"
    ]:

        # ----------------------------------------------------
        # 75%+ = broad systemic pattern
        # ----------------------------------------------------

        if prevalence >= 0.75:

            classification = (
                "SYSTEMATIQUE"
            )


        # ----------------------------------------------------
        # 15% to 75% = recurring pattern
        #
        # Changed from 25% because 4/20 = 20% is clearly more
        # useful to describe as recurring than isolated.
        # ----------------------------------------------------

        elif prevalence >= 0.15:

            classification = (
                "RECURRENT"
            )


        # ----------------------------------------------------
        # Below 15% = isolated
        # ----------------------------------------------------

        else:

            classification = (
                "ISOLE"
            )


        classifications.append(
            classification
        )


    patterns[
        "pattern_type"
    ] = classifications


    interpretations = []


    for _, row in patterns.iterrows():

        pattern_type = row[
            "pattern_type"
        ]


        if pattern_type == "SYSTEMATIQUE":

            interpretation = (
                "Écart observé sur une grande partie de la "
                "population. Vérifier en priorité une règle de "
                "transformation, une configuration globale ou "
                "une limite du jeu de données."
            )


        elif pattern_type == "RECURRENT":

            interpretation = (
                "Écart récurrent mais non généralisé. "
                "Rechercher une cause commune touchant un "
                "sous-ensemble d'employés."
            )


        else:

            interpretation = (
                "Écart relativement isolé. Une investigation "
                "individuelle est probablement plus pertinente."
            )


        interpretations.append(
            interpretation
        )


    patterns[
        "ai_interpretation"
    ] = interpretations


    return (
        patterns
        .sort_values(
            by=[
                "affected_employees",
                "occurrences",
            ],
            ascending=[
                False,
                False,
            ],
        )
        .reset_index(
            drop=True
        )
    )


# ============================================================
# TERMINAL PREVIEW
# ============================================================

def preview_ai_analysis(
    ambiguous_df,
    priority_df,
    pattern_df,
):
    """
    Display AI analysis.
    """

    print(
        "\n"
        +
        "=" * 70
    )

    print(
        "CORROBORIA - AI ANALYSIS"
    )

    print(
        "=" * 70
    )


    # ========================================================
    # AMBIGUOUS ASSIGNMENTS
    # ========================================================

    print(
        "\nAMBIGUOUS ASSIGNMENT ANALYSIS"
    )

    print(
        "-" * 70
    )


    if ambiguous_df.empty:

        print(
            "No ambiguous assignments analyzed."
        )

    else:

        columns = [
            "employee_id",
            "assignment_type",
            "source_row_index",
            "destination_row_index",
            "pair_match_score",
            "selected_by_best_matching",
            "best_matching_score",
            "second_matching_score",
            "matching_confidence",
            "recommendation",
        ]


        print(
            ambiguous_df[
                columns
            ]
            .to_string(
                index=False
            )
        )


    # ========================================================
    # PRIORITIZATION
    # ========================================================

    print(
        "\nAI PRIORITIZATION"
    )

    print(
        "-" * 70
    )


    if priority_df.empty:

        print(
            "No anomaly priorities generated."
        )

    else:

        columns = [
            "employee_id",
            "anomaly_count",
            "ai_anomaly_score",
            "ai_priority",
            "ai_explanation",
        ]


        print(
            priority_df[
                columns
            ]
            .head(
                20
            )
            .to_string(
                index=False
            )
        )


    # ========================================================
    # PATTERNS
    # ========================================================

    print(
        "\nDETECTED PATTERNS"
    )

    print(
        "-" * 70
    )


    if pattern_df.empty:

        print(
            "No discrepancy patterns detected."
        )

    else:

        columns = [
            "field",
            "verdict",
            "affected_employees",
            "employee_prevalence",
            "pattern_type",
        ]


        print(
            pattern_df[
                columns
            ]
            .to_string(
                index=False
            )
        )


# ============================================================
# EXCEL FORMATTING
# ============================================================

def auto_size_columns(
    worksheet,
    max_width=45,
):
    """
    Set readable Excel column widths.
    """

    for column_cells in worksheet.columns:

        maximum = 0


        for cell in column_cells:

            if cell.value is None:

                continue


            maximum = max(
                maximum,
                len(
                    str(
                        cell.value
                    )
                ),
            )


        letter = get_column_letter(
            column_cells[
                0
            ].column
        )


        worksheet.column_dimensions[
            letter
        ].width = min(
            max(
                maximum + 2,
                10,
            ),
            max_width,
        )


def style_sheet(
    worksheet,
):
    """
    Style AI worksheet.
    """

    worksheet.freeze_panes = (
        "A2"
    )


    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )


    fill = PatternFill(
        fill_type="solid",
        fgColor="7030A0",
    )


    font = Font(
        color="FFFFFF",
        bold=True,
    )


    for cell in worksheet[
        1
    ]:

        cell.fill = (
            fill
        )

        cell.font = (
            font
        )

        cell.alignment = Alignment(
            wrap_text=True,
            vertical="center",
        )


    for row in worksheet.iter_rows(
        min_row=2
    ):

        for cell in row:

            cell.alignment = Alignment(
                wrap_text=True,
                vertical="top",
            )


    auto_size_columns(
        worksheet
    )


# ============================================================
# EXPORT AI ANALYSIS
# ============================================================

def export_ai_analysis(
    ambiguous_df,
    priority_df,
    pattern_df,
    excel_path: Path,
    output_dir: Path,
):
    """
    Export AI outputs and append/replace AI sheets in main
    workbook.
    """

    output_dir.mkdir(
        parents=True,
        exist_ok=True,
    )


    ambiguous_csv = (
        output_dir
        /
        "corroboria_ai_assignment_analysis.csv"
    )


    priority_csv = (
        output_dir
        /
        "corroboria_ai_priorities.csv"
    )


    patterns_csv = (
        output_dir
        /
        "corroboria_ai_patterns.csv"
    )


    ambiguous_df.to_csv(
        ambiguous_csv,
        index=False,
        encoding="utf-8-sig",
    )


    priority_df.to_csv(
        priority_csv,
        index=False,
        encoding="utf-8-sig",
    )


    pattern_df.to_csv(
        patterns_csv,
        index=False,
        encoding="utf-8-sig",
    )


    # ========================================================
    # APPEND TO EXCEL REPORT
    # ========================================================

    with pd.ExcelWriter(
        excel_path,
        engine="openpyxl",
        mode="a",
        if_sheet_exists="replace",
    ) as writer:

        ambiguous_df.to_excel(
            writer,
            sheet_name="Analyse IA",
            index=False,
        )


        priority_df.to_excel(
            writer,
            sheet_name="Priorisation IA",
            index=False,
        )


        pattern_df.to_excel(
            writer,
            sheet_name="Patterns IA",
            index=False,
        )


        for sheet_name in [
            "Analyse IA",
            "Priorisation IA",
            "Patterns IA",
        ]:

            style_sheet(
                writer.sheets[
                    sheet_name
                ]
            )


    return {
        "assignment_analysis_csv":
            ambiguous_csv,

        "priorities_csv":
            priority_csv,

        "patterns_csv":
            patterns_csv,

        "excel":
            excel_path,
    }
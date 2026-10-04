# ============================================================
# CORROBORIA
# UNIFIED FINAL REPORT
# ============================================================
#
# PURPOSE
#
# Combine:
#
#   1. Direct field comparisons
#   2. Deterministic business-rule comparisons
#   3. Structural assignment problems
#
# into one standardized corroboration report.
#
#
# FINAL VERDICTS
#
#   CONFORME
#
#       No meaningful discrepancy.
#
#
#   ECART_JUSTIFIE
#
#       Raw values differ, but normalization or an accepted
#       transformation explains the difference.
#
#
#   ANOMALIE
#
#       Deterministic comparison/rule proves that System B
#       does not contain the expected value.
#
#
#   A_INVESTIGUER
#
#       A safe deterministic verdict cannot be reached.
#
#
# IMPORTANT
#
# AI does NOT override deterministic rules.
#
# Dataset limitations may remain A_INVESTIGUER without being
# sent to AI if AI cannot reasonably solve the missing
# information problem.
#
# ============================================================


# ============================================================
# IMPORTS
# ============================================================

from pathlib import Path
import json

import pandas as pd

from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)

from openpyxl.utils import (
    get_column_letter,
)


# ============================================================
# GENERAL HELPERS
# ============================================================

def safe_value(value):
    """
    Convert pandas missing values to None.
    """

    if value is None:
        return None


    try:

        if pd.isna(value):
            return None

    except (
        TypeError,
        ValueError,
    ):

        pass


    return value


def safe_text(value):
    """
    Convert optional values to text.
    """

    value = safe_value(
        value
    )


    if value is None:
        return None


    return str(
        value
    )


def serialize_inputs(value):
    """
    Serialize rule inputs for CSV / Excel traceability.
    """

    if value is None:
        return None


    if isinstance(
        value,
        dict,
    ):

        return json.dumps(
            value,
            ensure_ascii=False,
            default=str,
        )


    return str(
        value
    )


# ============================================================
# PRIORITY
# ============================================================

def assign_priority(
    verdict,
    comparison_type,
):
    """
    Deterministic baseline priority.

    AI may later enrich priority, but it cannot change the
    deterministic verdict.
    """

    if verdict == "ANOMALIE":

        if comparison_type in {
            "SOURCE_ONLY_ASSIGNMENT",
            "DESTINATION_ONLY_ASSIGNMENT",
        }:

            return "CRITIQUE"


        return "ÉLEVÉE"


    if verdict == "A_INVESTIGUER":

        if comparison_type == "AMBIGUOUS_ASSIGNMENT_MATCH":

            return "ÉLEVÉE"


        return "MOYENNE"


    return "FAIBLE"


# ============================================================
# CONFIDENCE
# ============================================================

def assign_confidence(
    verdict,
    decision_source,
):
    """
    Confidence in the verdict.

    This is NOT the probability that the employee data itself
    is correct.
    """

    if decision_source in {
        "DIRECT",
        "RULE",
        "STRUCTURAL",
    }:

        if verdict in {
            "CONFORME",
            "ECART_JUSTIFIE",
            "ANOMALIE",
        }:

            return 1.0


    if decision_source == "RULE_LIMITATION":

        return 0.5


    if verdict == "A_INVESTIGUER":

        return 0.5


    return None


# ============================================================
# DATASET LIMITATIONS
# ============================================================

def apply_dataset_limitations(
    rule_id,
    current_verdict,
    explanation,
):
    """
    Handle known cases where the anonymized challenge dataset
    prevents CorroborIA from safely deciding whether a
    discrepancy represents a real production error.

    EMAIL CASE
    ----------

    Mapping.xlsx gives a deterministic email construction rule.

    However, every System B email follows a separate systematic
    anonymized convention such as:

        dev-08-v2_...

    The challenge documentation confirms that the extracts are
    anonymized, but does not provide a transformation that
    allows that destination anonymization to be reconstructed
    from the anonymized source.

    Therefore:

        email mismatch
            -> A_INVESTIGUER

    BUT:

        requires_ai = False

    AI cannot reconstruct information removed by anonymization.
    """

    if (
        rule_id == "RULE_EMAIL"
        and
        current_verdict == "ANOMALIE"
    ):

        return {
            "verdict":
                "A_INVESTIGUER",

            "decision_source":
                "RULE_LIMITATION",

            "requires_ai":
                False,

            "limitation_code":
                "ANONYMIZATION_EMAIL",

            "explanation":
                (
                    "La règle métier de construction du courriel "
                    "a été appliquée correctement. Toutefois, les "
                    "courriels du système B suivent une convention "
                    "systématique différente qui semble liée à "
                    "l'anonymisation du jeu de données. Aucune règle "
                    "fournie ne permet de reconstruire cette "
                    "transformation. L'écart est donc conservé comme "
                    "cas à investiguer, mais n'est pas envoyé à l'IA "
                    "puisque l'information nécessaire a été supprimée "
                    "par l'anonymisation. "
                    f"Détail déterministe initial : {explanation}"
                ),
        }


    return {
        "verdict":
            current_verdict,

        "decision_source":
            "RULE",

        "requires_ai":
            current_verdict
            ==
            "A_INVESTIGUER",

        "limitation_code":
            None,

        "explanation":
            explanation,
    }


# ============================================================
# DIRECT COMPARISONS
# ============================================================

def convert_direct_results(
    comparison_df: pd.DataFrame,
) -> list[dict]:
    """
    Convert direct comparison results to unified schema.
    """

    records = []


    for _, row in comparison_df.iterrows():

        comparison_status = (
            row[
                "comparison_status"
            ]
        )


        # ====================================================
        # FINAL DIRECT VERDICT
        # ====================================================

        if comparison_status == "EXACT_MATCH":

            verdict = (
                "CONFORME"
            )

            explanation = (
                "La valeur du système A correspond exactement "
                "à la valeur du système B."
            )


        elif comparison_status == "NORMALIZED_MATCH":

            verdict = (
                "ECART_JUSTIFIE"
            )

            explanation = (
                "Les valeurs brutes diffèrent, mais deviennent "
                "équivalentes après normalisation. L'écart est "
                "justifié par une différence de représentation."
            )


        else:

            verdict = (
                "ANOMALIE"
            )

            explanation = (
                "Le champ possède une correspondance directe "
                "dans le mapping et les valeurs normalisées "
                "demeurent différentes. Aucune règle métier "
                "fournie ne justifie cet écart."
            )


        decision_source = (
            "DIRECT"
        )


        comparison_type = (
            comparison_status
        )


        records.append(
            {
                "employee_id":
                    safe_value(
                        row.get(
                            "employee_id"
                        )
                    ),

                "assignment_type":
                    safe_value(
                        row.get(
                            "assignment_type"
                        )
                    ),

                "source_row_index":
                    safe_value(
                        row.get(
                            "source_row_index"
                        )
                    ),

                "destination_row_index":
                    safe_value(
                        row.get(
                            "destination_row_index"
                        )
                    ),

                "field":
                    safe_value(
                        row.get(
                            "destination_field"
                        )
                    ),

                "source_field":
                    safe_value(
                        row.get(
                            "source_field"
                        )
                    ),

                "source_raw_value":
                    safe_value(
                        row.get(
                            "source_raw_value"
                        )
                    ),

                "source_normalized_value":
                    safe_value(
                        row.get(
                            "source_normalized_value"
                        )
                    ),

                "expected_value":
                    safe_value(
                        row.get(
                            "source_normalized_value"
                        )
                    ),

                "destination_raw_value":
                    safe_value(
                        row.get(
                            "destination_raw_value"
                        )
                    ),

                "destination_normalized_value":
                    safe_value(
                        row.get(
                            "destination_normalized_value"
                        )
                    ),

                "comparison_layer":
                    "DIRECT",

                "comparison_type":
                    comparison_type,

                "mapping_row":
                    safe_value(
                        row.get(
                            "mapping_row"
                        )
                    ),

                "rule_id":
                    None,

                "rule_name":
                    None,

                "rule_inputs":
                    None,

                "decision_source":
                    decision_source,

                "limitation_code":
                    None,

                "verdict":
                    verdict,

                "priority":
                    assign_priority(
                        verdict=
                            verdict,

                        comparison_type=
                            comparison_type,
                    ),

                "confidence":
                    assign_confidence(
                        verdict=
                            verdict,

                        decision_source=
                            decision_source,
                    ),

                "explanation":
                    explanation,

                "requires_ai":
                    False,

                "is_investigation_item":
                    verdict
                    in {
                        "ANOMALIE",
                        "A_INVESTIGUER",
                    },
            }
        )


    return records


# ============================================================
# BUSINESS RULE RESULTS
# ============================================================

def convert_rule_results(
    rule_results_df: pd.DataFrame,
) -> list[dict]:
    """
    Convert deterministic business-rule results into the
    unified schema.
    """

    records = []


    for _, row in rule_results_df.iterrows():

        current_verdict = (
            row[
                "verdict"
            ]
        )


        comparison_status = (
            row[
                "comparison_status"
            ]
        )


        rule_id = (
            row[
                "rule_id"
            ]
        )


        original_explanation = (
            safe_text(
                row.get(
                    "explanation"
                )
            )
            or
            ""
        )


        # ====================================================
        # APPLY KNOWN DATASET LIMITATIONS
        # ====================================================

        limitation_result = (
            apply_dataset_limitations(

                rule_id=
                    rule_id,

                current_verdict=
                    current_verdict,

                explanation=
                    original_explanation,
            )
        )


        verdict = (
            limitation_result[
                "verdict"
            ]
        )


        decision_source = (
            limitation_result[
                "decision_source"
            ]
        )


        explanation = (
            limitation_result[
                "explanation"
            ]
        )


        requires_ai = (
            limitation_result[
                "requires_ai"
            ]
        )


        limitation_code = (
            limitation_result[
                "limitation_code"
            ]
        )


        comparison_type = (
            comparison_status
        )


        records.append(
            {
                "employee_id":
                    safe_value(
                        row.get(
                            "employee_id"
                        )
                    ),

                "assignment_type":
                    safe_value(
                        row.get(
                            "assignment_type"
                        )
                    ),

                "source_row_index":
                    safe_value(
                        row.get(
                            "source_row_index"
                        )
                    ),

                "destination_row_index":
                    safe_value(
                        row.get(
                            "destination_row_index"
                        )
                    ),

                "field":
                    safe_value(
                        row.get(
                            "destination_field"
                        )
                    ),

                "source_field":
                    None,

                "source_raw_value":
                    None,

                "source_normalized_value":
                    None,

                "expected_value":
                    safe_value(
                        row.get(
                            "expected_value"
                        )
                    ),

                "destination_raw_value":
                    safe_value(
                        row.get(
                            "destination_raw_value"
                        )
                    ),

                "destination_normalized_value":
                    safe_value(
                        row.get(
                            "destination_normalized_value"
                        )
                    ),

                "comparison_layer":
                    "BUSINESS_RULE",

                "comparison_type":
                    comparison_type,

                "mapping_row":
                    safe_value(
                        row.get(
                            "mapping_row"
                        )
                    ),

                "rule_id":
                    rule_id,

                "rule_name":
                    safe_value(
                        row.get(
                            "rule_name"
                        )
                    ),

                "rule_inputs":
                    serialize_inputs(
                        row.get(
                            "source_inputs"
                        )
                    ),

                "decision_source":
                    decision_source,

                "limitation_code":
                    limitation_code,

                "verdict":
                    verdict,

                "priority":
                    assign_priority(
                        verdict=
                            verdict,

                        comparison_type=
                            comparison_type,
                    ),

                "confidence":
                    assign_confidence(
                        verdict=
                            verdict,

                        decision_source=
                            decision_source,
                    ),

                "explanation":
                    explanation,

                "requires_ai":
                    requires_ai,

                "is_investigation_item":
                    verdict
                    in {
                        "ANOMALIE",
                        "A_INVESTIGUER",
                    },
            }
        )


    return records


# ============================================================
# STRUCTURAL ASSIGNMENT RESULTS
# ============================================================

def convert_assignment_results(
    assignment_matches_df: pd.DataFrame,
) -> list[dict]:
    """
    Add assignment-level structural problems.

    MATCHED assignments are already represented by direct and
    business-rule results, so they are not duplicated here.
    """

    records = []


    for _, row in assignment_matches_df.iterrows():

        status = (
            row[
                "assignment_match_status"
            ]
        )


        if status == "MATCHED":

            continue


        employee_id = safe_value(
            row.get(
                "employee_id"
            )
        )


        assignment_type = safe_value(
            row.get(
                "assignment_type"
            )
        )


        source_index = safe_value(
            row.get(
                "source_row_index"
            )
        )


        destination_index = safe_value(
            row.get(
                "destination_row_index"
            )
        )


        source_candidates = safe_value(
            row.get(
                "source_candidate_indices"
            )
        )


        destination_candidates = safe_value(
            row.get(
                "destination_candidate_indices"
            )
        )


        # ====================================================
        # SOURCE-ONLY
        # ====================================================

        if status == "SOURCE_ONLY_ASSIGNMENT":

            verdict = (
                "ANOMALIE"
            )

            comparison_type = (
                "SOURCE_ONLY_ASSIGNMENT"
            )

            explanation = (
                "Une affectation existe dans le système A, "
                "mais aucune affectation correspondante du "
                "même type n'a été trouvée dans le système B."
            )

            requires_ai = (
                False
            )


        # ====================================================
        # DESTINATION-ONLY
        # ====================================================

        elif status == "DESTINATION_ONLY_ASSIGNMENT":

            verdict = (
                "ANOMALIE"
            )

            comparison_type = (
                "DESTINATION_ONLY_ASSIGNMENT"
            )

            explanation = (
                "Une affectation existe dans le système B, "
                "mais aucune affectation correspondante du "
                "même type n'a été trouvée dans le système A."
            )

            requires_ai = (
                False
            )


        # ====================================================
        # AMBIGUOUS MATCH
        # ====================================================

        elif status == "AMBIGUOUS_ASSIGNMENT_MATCH":

            verdict = (
                "A_INVESTIGUER"
            )

            comparison_type = (
                "AMBIGUOUS_ASSIGNMENT_MATCH"
            )

            explanation = (
                "Plusieurs affectations candidates du même "
                "type existent dans les deux systèmes. Une "
                "correspondance unique ne peut pas être "
                "déterminée par les règles déterministes."
            )

            requires_ai = (
                True
            )


        # ====================================================
        # UNKNOWN TYPE
        # ====================================================

        elif status == "UNRESOLVED_ASSIGNMENT_TYPE":

            verdict = (
                "A_INVESTIGUER"
            )

            comparison_type = (
                "UNRESOLVED_ASSIGNMENT_TYPE"
            )

            explanation = (
                "Le type d'affectation ne peut pas être "
                "déterminé à partir des valeurs fournies."
            )

            requires_ai = (
                True
            )


        else:

            verdict = (
                "A_INVESTIGUER"
            )

            comparison_type = (
                status
            )

            explanation = (
                f"Situation structurelle non reconnue : "
                f"{status}."
            )

            requires_ai = (
                True
            )


        decision_source = (
            "STRUCTURAL"
        )


        records.append(
            {
                "employee_id":
                    employee_id,

                "assignment_type":
                    assignment_type,

                "source_row_index":
                    source_index,

                "destination_row_index":
                    destination_index,

                "field":
                    "__ASSIGNMENT__",

                "source_field":
                    None,

                "source_raw_value":
                    None,

                "source_normalized_value":
                    None,

                "expected_value":
                    None,

                "destination_raw_value":
                    None,

                "destination_normalized_value":
                    None,

                "comparison_layer":
                    "STRUCTURAL",

                "comparison_type":
                    comparison_type,

                "mapping_row":
                    None,

                "rule_id":
                    None,

                "rule_name":
                    None,

                "rule_inputs":
                    json.dumps(
                        {
                            "source_candidates":
                                source_candidates,

                            "destination_candidates":
                                destination_candidates,
                        },
                        ensure_ascii=False,
                        default=str,
                    ),

                "decision_source":
                    decision_source,

                "limitation_code":
                    None,

                "verdict":
                    verdict,

                "priority":
                    assign_priority(
                        verdict=
                            verdict,

                        comparison_type=
                            comparison_type,
                    ),

                "confidence":
                    assign_confidence(
                        verdict=
                            verdict,

                        decision_source=
                            decision_source,
                    ),

                "explanation":
                    explanation,

                "requires_ai":
                    requires_ai,

                "is_investigation_item":
                    True,
            }
        )


    return records


# ============================================================
# BUILD COMPLETE REPORT
# ============================================================

def build_final_report(
    comparison_df: pd.DataFrame,
    rule_results_df: pd.DataFrame,
    assignment_matches_df: pd.DataFrame,
) -> pd.DataFrame:
    """
    Build the complete CorroborIA trace.
    """

    records = []


    records.extend(
        convert_direct_results(
            comparison_df=
                comparison_df
        )
    )


    records.extend(
        convert_rule_results(
            rule_results_df=
                rule_results_df
        )
    )


    records.extend(
        convert_assignment_results(
            assignment_matches_df=
                assignment_matches_df
        )
    )


    report_df = pd.DataFrame(
        records
    )


    if report_df.empty:

        return report_df


    # ========================================================
    # SORT
    # ========================================================

    verdict_order = {
        "ANOMALIE": 0,
        "A_INVESTIGUER": 1,
        "ECART_JUSTIFIE": 2,
        "CONFORME": 3,
    }


    priority_order = {
        "CRITIQUE": 0,
        "ÉLEVÉE": 1,
        "MOYENNE": 2,
        "FAIBLE": 3,
    }


    report_df[
        "_verdict_order"
    ] = report_df[
        "verdict"
    ].map(
        verdict_order
    )


    report_df[
        "_priority_order"
    ] = report_df[
        "priority"
    ].map(
        priority_order
    )


    report_df = (
        report_df
        .sort_values(
            by=[
                "_verdict_order",
                "_priority_order",
                "employee_id",
                "field",
            ],
            na_position="last",
        )
        .drop(
            columns=[
                "_verdict_order",
                "_priority_order",
            ]
        )
        .reset_index(
            drop=True
        )
    )


    return report_df


# ============================================================
# INVESTIGATION VIEW
# ============================================================

def build_investigation_report(
    final_report_df: pd.DataFrame,
) -> pd.DataFrame:
    """
    Technical investigation subset.
    """

    if final_report_df.empty:

        return final_report_df.copy()


    return (
        final_report_df[
            final_report_df[
                "verdict"
            ].isin(
                [
                    "ANOMALIE",
                    "A_INVESTIGUER",
                ]
            )
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )


# ============================================================
# AI QUEUE
# ============================================================

def build_ai_queue(
    final_report_df: pd.DataFrame,
) -> pd.DataFrame:
    """
    Keep ONLY cases for which AI can actually add value.

    Email anonymization limitations are therefore excluded.
    """

    if final_report_df.empty:

        return final_report_df.copy()


    return (
        final_report_df[
            final_report_df[
                "requires_ai"
            ]
            ==
            True
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )


# ============================================================
# ANALYST-FRIENDLY INVESTIGATION VIEW
# ============================================================

def build_analyst_investigation_view(
    investigation_df: pd.DataFrame,
) -> pd.DataFrame:
    """
    Create the simplified sheet intended for functional users.

    The full technical trace remains available elsewhere.
    """

    columns = [
        "employee_id",
        "assignment_type",
        "field",
        "expected_value",
        "destination_normalized_value",
        "verdict",
        "priority",
        "rule_name",
        "decision_source",
        "limitation_code",
        "confidence",
        "explanation",
    ]


    available_columns = [
        column
        for column in columns
        if column in investigation_df.columns
    ]


    result = (
        investigation_df[
            available_columns
        ]
        .copy()
    )


    rename_map = {
        "employee_id":
            "Employé",

        "assignment_type":
            "Affectation",

        "field":
            "Champ",

        "expected_value":
            "Valeur attendue",

        "destination_normalized_value":
            "Valeur reçue",

        "verdict":
            "Verdict",

        "priority":
            "Priorité",

        "rule_name":
            "Règle appliquée",

        "decision_source":
            "Source décision",

        "limitation_code":
            "Limite connue",

        "confidence":
            "Confiance",

        "explanation":
            "Explication",
    }


    result = result.rename(
        columns=
            rename_map
    )


    return result


# ============================================================
# JUSTIFIED DISCREPANCY VIEW
# ============================================================

def build_justified_view(
    final_report_df: pd.DataFrame,
) -> pd.DataFrame:
    """
    User-facing list of automatically justified differences.
    """

    justified = (
        final_report_df[
            final_report_df[
                "verdict"
            ]
            ==
            "ECART_JUSTIFIE"
        ]
        .copy()
    )


    columns = [
        "employee_id",
        "assignment_type",
        "field",
        "source_raw_value",
        "destination_raw_value",
        "source_normalized_value",
        "destination_normalized_value",
        "explanation",
    ]


    available_columns = [
        column
        for column in columns
        if column in justified.columns
    ]


    result = justified[
        available_columns
    ].copy()


    result = result.rename(
        columns={
            "employee_id":
                "Employé",

            "assignment_type":
                "Affectation",

            "field":
                "Champ",

            "source_raw_value":
                "Valeur A brute",

            "destination_raw_value":
                "Valeur B brute",

            "source_normalized_value":
                "Valeur A normalisée",

            "destination_normalized_value":
                "Valeur B normalisée",

            "explanation":
                "Justification",
        }
    )


    return result


# ============================================================
# SUMMARY DATA
# ============================================================

def build_summary_tables(
    final_report_df: pd.DataFrame,
    investigation_df: pd.DataFrame,
    ai_queue_df: pd.DataFrame,
):
    """
    Build dashboard-style summary tables for the workbook.
    """

    verdict_counts = (
        final_report_df[
            "verdict"
        ]
        .value_counts()
    )


    summary_df = pd.DataFrame(
        [
            {
                "Indicateur":
                    "Comparaisons totales",

                "Valeur":
                    len(
                        final_report_df
                    ),
            },
            {
                "Indicateur":
                    "Conformes",

                "Valeur":
                    int(
                        verdict_counts.get(
                            "CONFORME",
                            0,
                        )
                    ),
            },
            {
                "Indicateur":
                    "Écarts justifiés",

                "Valeur":
                    int(
                        verdict_counts.get(
                            "ECART_JUSTIFIE",
                            0,
                        )
                    ),
            },
            {
                "Indicateur":
                    "Anomalies",

                "Valeur":
                    int(
                        verdict_counts.get(
                            "ANOMALIE",
                            0,
                        )
                    ),
            },
            {
                "Indicateur":
                    "À investiguer",

                "Valeur":
                    int(
                        verdict_counts.get(
                            "A_INVESTIGUER",
                            0,
                        )
                    ),
            },
            {
                "Indicateur":
                    "Éléments nécessitant attention",

                "Valeur":
                    len(
                        investigation_df
                    ),
            },
            {
                "Indicateur":
                    "Cas réellement envoyés à l'IA",

                "Valeur":
                    len(
                        ai_queue_df
                    ),
            },
        ]
    )


    # ========================================================
    # LAYER SUMMARY
    # ========================================================

    layer_df = (
        pd.crosstab(
            final_report_df[
                "comparison_layer"
            ],
            final_report_df[
                "verdict"
            ],
        )
        .reset_index()
    )


    # ========================================================
    # ISSUE SUMMARY
    # ========================================================

    if investigation_df.empty:

        issue_df = pd.DataFrame(
            columns=[
                "Champ",
                "Verdict",
                "Nombre",
            ]
        )


    else:

        issue_df = (
            investigation_df
            .groupby(
                [
                    "field",
                    "verdict",
                ],
                dropna=False,
            )
            .size()
            .reset_index(
                name="Nombre"
            )
            .rename(
                columns={
                    "field":
                        "Champ",

                    "verdict":
                        "Verdict",
                }
            )
            .sort_values(
                by="Nombre",
                ascending=False,
            )
            .reset_index(
                drop=True
            )
        )


    return (
        summary_df,
        layer_df,
        issue_df,
    )


# ============================================================
# PREVIEW FINAL REPORT
# ============================================================

def preview_final_report(
    final_report_df: pd.DataFrame,
    investigation_df: pd.DataFrame,
    ai_queue_df: pd.DataFrame,
) -> None:
    """
    Terminal summary.
    """

    print(
        "\n"
        +
        "=" * 70
    )

    print(
        "CORROBORIA - UNIFIED CORROBORATION REPORT"
    )

    print(
        "=" * 70
    )


    if final_report_df.empty:

        print(
            "\nNo report rows were produced."
        )

        return


    print(
        f"\nTotal report rows: "
        f"{len(final_report_df)}"
    )


    # ========================================================
    # VERDICTS
    # ========================================================

    print(
        "\nFINAL VERDICTS"
    )

    print(
        "-" * 70
    )


    for (
        verdict,
        count,
    ) in (
        final_report_df[
            "verdict"
        ]
        .value_counts()
        .items()
    ):

        print(
            f"{verdict}: {count}"
        )


    # ========================================================
    # LAYERS
    # ========================================================

    print(
        "\nRESULTS BY LAYER"
    )

    print(
        "-" * 70
    )


    print(
        pd.crosstab(
            final_report_df[
                "comparison_layer"
            ],
            final_report_df[
                "verdict"
            ],
        ).to_string()
    )


    # ========================================================
    # ATTENTION
    # ========================================================

    print(
        "\nITEMS REQUIRING ATTENTION"
    )

    print(
        "-" * 70
    )


    print(
        f"Total investigation items: "
        f"{len(investigation_df)}"
    )


    # ========================================================
    # AI
    # ========================================================

    print(
        "\nCASES ELIGIBLE FOR AI ANALYSIS"
    )

    print(
        "-" * 70
    )


    print(
        f"AI queue rows: "
        f"{len(ai_queue_df)}"
    )


    if ai_queue_df.empty:

        print(
            "None."
        )

    else:

        preview_columns = [
            "employee_id",
            "assignment_type",
            "field",
            "comparison_type",
            "verdict",
            "explanation",
        ]


        print(
            ai_queue_df[
                preview_columns
            ]
            .to_string(
                index=False
            )
        )


# ============================================================
# EXCEL FORMATTING HELPERS
# ============================================================

def auto_size_columns(
    worksheet,
    max_width=45,
):
    """
    Adjust Excel column widths.
    """

    for column_cells in worksheet.columns:

        maximum_length = 0


        for cell in column_cells:

            value = cell.value


            if value is None:

                continue


            length = len(
                str(
                    value
                )
            )


            maximum_length = max(
                maximum_length,
                length,
            )


        width = min(
            maximum_length + 2,
            max_width,
        )


        column_letter = (
            get_column_letter(
                column_cells[
                    0
                ].column
            )
        )


        worksheet.column_dimensions[
            column_letter
        ].width = max(
            width,
            10,
        )


def style_table_sheet(
    worksheet,
    freeze="A2",
):
    """
    Apply basic analyst-friendly formatting.
    """

    worksheet.freeze_panes = (
        freeze
    )


    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )


    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )


    header_font = Font(
        color="FFFFFF",
        bold=True,
    )


    for cell in worksheet[
        1
    ]:

        cell.fill = (
            header_fill
        )

        cell.font = (
            header_font
        )

        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )


    for row in worksheet.iter_rows(
        min_row=2
    ):

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )


    auto_size_columns(
        worksheet
    )


def style_summary_sheet(
    worksheet,
):
    """
    Format Résumé sheet.
    """

    worksheet.freeze_panes = (
        "A2"
    )


    title_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )


    title_font = Font(
        color="FFFFFF",
        bold=True,
    )


    for cell in worksheet[
        1
    ]:

        if cell.value is not None:

            cell.fill = (
                title_fill
            )

            cell.font = (
                title_font
            )


    auto_size_columns(
        worksheet,
        max_width=35,
    )


# ============================================================
# EXPORT REPORTS
# ============================================================

def export_reports(
    final_report_df: pd.DataFrame,
    investigation_df: pd.DataFrame,
    ai_queue_df: pd.DataFrame,
    output_dir: Path,
) -> dict[str, Path]:
    """
    Export CSV + polished Excel workbook.
    """

    output_dir.mkdir(
        parents=True,
        exist_ok=True,
    )


    # ========================================================
    # PATHS
    # ========================================================

    full_csv = (
        output_dir
        /
        "corroboria_full_report.csv"
    )


    investigation_csv = (
        output_dir
        /
        "corroboria_investigation_report.csv"
    )


    ai_csv = (
        output_dir
        /
        "corroboria_ai_queue.csv"
    )


    excel_path = (
        output_dir
        /
        "corroboria_report.xlsx"
    )


    # ========================================================
    # USER-FACING VIEWS
    # ========================================================

    analyst_investigation_df = (
        build_analyst_investigation_view(
            investigation_df=
                investigation_df
        )
    )


    justified_df = (
        build_justified_view(
            final_report_df=
                final_report_df
        )
    )


    (
        summary_df,
        layer_df,
        issue_df,
    ) = build_summary_tables(

        final_report_df=
            final_report_df,

        investigation_df=
            investigation_df,

        ai_queue_df=
            ai_queue_df,
    )


    # ========================================================
    # CSV EXPORTS
    # ========================================================

    final_report_df.to_csv(
        full_csv,
        index=False,
        encoding="utf-8-sig",
    )


    analyst_investigation_df.to_csv(
        investigation_csv,
        index=False,
        encoding="utf-8-sig",
    )


    ai_queue_df.to_csv(
        ai_csv,
        index=False,
        encoding="utf-8-sig",
    )


    # ========================================================
    # EXCEL EXPORT
    # ========================================================

    with pd.ExcelWriter(
        excel_path,
        engine="openpyxl",
    ) as writer:

        # ====================================================
        # SHEET 1
        # RÉSUMÉ
        # ====================================================

        summary_df.to_excel(
            writer,
            sheet_name="Résumé",
            index=False,
            startrow=1,
            startcol=0,
        )


        layer_df.to_excel(
            writer,
            sheet_name="Résumé",
            index=False,
            startrow=1,
            startcol=4,
        )


        issue_df.to_excel(
            writer,
            sheet_name="Résumé",
            index=False,
            startrow=11,
            startcol=0,
        )


        summary_sheet = (
            writer.sheets[
                "Résumé"
            ]
        )


        summary_sheet[
            "A1"
        ] = (
            "Résumé de corroboration"
        )


        summary_sheet[
            "E1"
        ] = (
            "Résultats par couche"
        )


        summary_sheet[
            "A11"
        ] = (
            "Principaux écarts"
        )


        # ====================================================
        # SHEET 2
        # ANALYST INVESTIGATION
        # ====================================================

        analyst_investigation_df.to_excel(
            writer,
            sheet_name="À investiguer",
            index=False,
        )


        # ====================================================
        # SHEET 3
        # JUSTIFIED DIFFERENCES
        # ====================================================

        justified_df.to_excel(
            writer,
            sheet_name="Écarts justifiés",
            index=False,
        )


        # ====================================================
        # SHEET 4
        # FULL TECHNICAL TRACE
        # ====================================================

        final_report_df.to_excel(
            writer,
            sheet_name="Rapport complet",
            index=False,
        )


        # ====================================================
        # SHEET 5
        # AI QUEUE
        # ====================================================

        ai_queue_df.to_excel(
            writer,
            sheet_name="File IA",
            index=False,
        )


        # ====================================================
        # FORMATTING
        # ====================================================

        style_summary_sheet(
            summary_sheet
        )


        for sheet_name in [
            "À investiguer",
            "Écarts justifiés",
            "Rapport complet",
            "File IA",
        ]:

            worksheet = (
                writer.sheets[
                    sheet_name
                ]
            )


            style_table_sheet(
                worksheet
            )


    return {
        "full_csv":
            full_csv,

        "investigation_csv":
            investigation_csv,

        "ai_csv":
            ai_csv,

        "excel":
            excel_path,
    }
"""End-to-end QA for CorroborIA."""

from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path

from openpyxl import load_workbook

from corroboria.export import (
    audit_to_json,
    cases_to_csv,
    run_to_excel,
)
from corroboria.pipeline import (
    InputPaths,
    run_reconciliation,
)


BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"


EXPECTED_COUNTS = {
    "Actual anomaly": 55,
    "Needs review": 21,
    "Justified difference": 20,
    "Match": 406,
}


class NamedBytesIO(BytesIO):
    """In-memory file that behaves like an uploaded file."""

    def __init__(
        self,
        data: bytes,
        name: str,
    ) -> None:

        super().__init__(
            data
        )

        self.name = name


def assert_counts(
    run,
) -> None:

    actual = (
        run.summary
        .set_index(
            "verdict"
        )[
            "field_comparisons_or_cases"
        ]
        .astype(
            int
        )
        .to_dict()
    )

    assert (
        actual
        ==
        EXPECTED_COUNTS
    ), (
        "\nUnexpected verdict counts.\n\n"
        f"Expected:\n{EXPECTED_COUNTS}\n\n"
        f"Actual:\n{actual}"
    )

    assert (
        len(
            run.final_report
        )
        ==
        502
    )

    assert (
        len(
            run.cases
        )
        ==
        502
    )


def uploaded_copy(
    path: Path,
) -> NamedBytesIO:

    return NamedBytesIO(
        path.read_bytes(),
        path.name,
    )


def build_uploaded_inputs(
    paths: InputPaths,
) -> InputPaths:

    return InputPaths(
        source=uploaded_copy(
            paths.source
        ),
        destination=uploaded_copy(
            paths.destination
        ),
        mapping=uploaded_copy(
            paths.mapping
        ),
        job_details=uploaded_copy(
            paths.job_details
        ),
        employment_reasons=uploaded_copy(
            paths.employment_reasons
        ),
    )


def assert_ai(
    run,
) -> None:

    assert (
        len(
            run.ai_queue
        )
        ==
        1
    ), (
        "Expected exactly one true AI queue case."
    )

    assert (
        len(
            run.ai_assignment_analysis
        )
        ==
        4
    ), (
        "Expected four candidate assignment pairs."
    )

    selected = (
        run.ai_assignment_analysis.loc[
            run.ai_assignment_analysis[
                "selected_by_best_matching"
            ].eq(
                True
            )
        ]
    )

    selected_pairs = {
        (
            int(
                row.source_row_index
            ),
            int(
                row.destination_row_index
            ),
        )
        for row
        in selected.itertuples()
    }

    expected_pairs = {
        (
            21,
            21,
        ),
        (
            22,
            19,
        ),
    }

    assert (
        selected_pairs
        ==
        expected_pairs
    ), (
        "Unexpected AI assignment proposal.\n"
        f"Expected: {expected_pairs}\n"
        f"Actual:   {selected_pairs}"
    )

    confidence = float(
        run.ai_assignment_analysis.iloc[
            0
        ][
            "matching_confidence"
        ]
    )

    assert (
        confidence
        >
        0.90
    ), (
        "Unexpectedly low matching confidence: "
        f"{confidence}"
    )

    assert (
        len(
            run.ai_patterns
        )
        ==
        9
    )


def assert_structural_case(
    run,
) -> None:

    structural = (
        run.cases.loc[
            run.cases[
                "record_identifier"
            ]
            .astype(
                str
            )
            .eq(
                "1545850"
            )
            &
            run.cases[
                "field"
            ]
            .astype(
                str
            )
            .eq(
                "__ASSIGNMENT__"
            )
        ]
    )

    assert (
        len(
            structural
        )
        ==
        1
    ), (
        "Expected one structural assignment "
        "case for employee 1545850."
    )

    row = structural.iloc[
        0
    ]

    assert (
        row[
            "verdict"
        ]
        ==
        "Actual anomaly"
    )

    assert (
        row[
            "assignment_type"
        ]
        ==
        "A"
    )


def assert_exports(
    run,
) -> None:

    # --------------------------------------------------------
    # CSV
    # --------------------------------------------------------

    csv_bytes = cases_to_csv(
        run.cases
    )

    assert (
        len(
            csv_bytes
        )
        >
        100
    )

    assert csv_bytes.startswith(
        b"\xef\xbb\xbf"
    )

    # --------------------------------------------------------
    # JSON AUDIT
    # --------------------------------------------------------

    audit_bytes = audit_to_json(
        run
    )

    audit = json.loads(
        audit_bytes.decode(
            "utf-8"
        )
    )

    assert (
        audit[
            "report_row_count"
        ]
        ==
        502
    )

    assert (
        audit[
            "ai_queue_rows"
        ]
        ==
        1
    )

    # --------------------------------------------------------
    # EXCEL
    # --------------------------------------------------------

    excel_bytes = run_to_excel(
        run
    )

    assert (
        len(
            excel_bytes
        )
        >
        1000
    )

    workbook = load_workbook(
        BytesIO(
            excel_bytes
        ),
        read_only=True,
    )

    expected_sheets = {
        "Summary",
        "Cases",
        "Investigation",
        "Full report",
        "Mapping",
        "Assignment matches",
        "Employee matches",
        "AI queue",
        "AI assignments",
        "AI priorities",
        "AI patterns",
        "Run audit",
    }

    missing = (
        expected_sheets
        -
        set(
            workbook.sheetnames
        )
    )

    assert (
        not missing
    ), (
        "Excel export is missing sheet(s): "
        f"{sorted(missing)}"
    )


def main() -> None:

    print(
        "=" * 72
    )

    print(
        "CORROBORIA - END-TO-END QA"
    )

    print(
        "=" * 72
    )

    paths = (
        InputPaths
        .from_data_directory(
            DATA_DIR
        )
    )

    # ========================================================
    # TEST 1
    # ========================================================

    print(
        "\n[1/6] Bundled data..."
    )

    bundled_run = (
        run_reconciliation(
            paths
        )
    )

    assert_counts(
        bundled_run
    )

    print(
        "PASS"
    )

    # ========================================================
    # TEST 2
    # ========================================================

    print(
        "\n[2/6] Structural anomaly..."
    )

    assert_structural_case(
        bundled_run
    )

    print(
        "PASS"
    )

    # ========================================================
    # TEST 3
    # ========================================================

    print(
        "\n[3/6] AI outputs..."
    )

    assert_ai(
        bundled_run
    )

    print(
        "PASS"
    )

    # ========================================================
    # TEST 4
    # ========================================================

    print(
        "\n[4/6] Export files..."
    )

    assert_exports(
        bundled_run
    )

    print(
        "PASS"
    )

    # ========================================================
    # TEST 5
    # ========================================================

    print(
        "\n[5/6] Simulated file uploads..."
    )

    upload_inputs = (
        build_uploaded_inputs(
            paths
        )
    )

    upload_run = (
        run_reconciliation(
            upload_inputs
        )
    )

    assert_counts(
        upload_run
    )

    print(
        "PASS"
    )

    # ========================================================
    # TEST 6
    # ========================================================

    print(
        "\n[6/6] Bundled vs uploaded results..."
    )

    bundled_counts = (
        bundled_run.summary
        .set_index(
            "verdict"
        )[
            "field_comparisons_or_cases"
        ]
        .to_dict()
    )

    uploaded_counts = (
        upload_run.summary
        .set_index(
            "verdict"
        )[
            "field_comparisons_or_cases"
        ]
        .to_dict()
    )

    assert (
        bundled_counts
        ==
        uploaded_counts
    )

    print(
        "PASS"
    )

    # ========================================================
    # SUCCESS
    # ========================================================

    print(
        "\n"
        +
        "=" * 72
    )

    print(
        "ALL QA CHECKS PASSED"
    )

    print(
        "=" * 72
    )

    print(
        "\nVerified baseline:"
    )

    print(
        "  Match:                 406"
    )

    print(
        "  Justified difference:   20"
    )

    print(
        "  Actual anomaly:          55"
    )

    print(
        "  Needs review:            21"
    )

    print(
        "  Total:                  502"
    )


if __name__ == "__main__":
    main()